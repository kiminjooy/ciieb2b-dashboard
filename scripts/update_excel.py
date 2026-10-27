#!/usr/bin/env python3
"""Manual GitHub Actions updater. Outputs public, minimal ciie-data.json."""
import json,os,re,sys,urllib.request
from io import BytesIO
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import load_workbook

URL=os.environ.get('ONEDRIVE_DOWNLOAD_URL','').strip()
if not URL: sys.exit('Set ONEDRIVE_DOWNLOAD_URL GitHub Actions variable')
if 'download=1' not in URL: URL+=('&' if '?' in URL else '?')+'download=1'
req=urllib.request.Request(URL,headers={'User-Agent':'Mozilla/5.0','Accept':'application/octet-stream,*/*'})
with urllib.request.urlopen(req,timeout=45) as resp:
 content=resp.read(30_000_000)
if not content.startswith(b'PK'):
 sys.exit('OneDrive did not return XLSX ZIP bytes (maybe HTML or authentication).')
wb=load_workbook(BytesIO(content),read_only=True,data_only=True)
for name in ('바이어','국내기업','Sheet1'):
 if name not in wb.sheetnames:sys.exit(f'Missing worksheet: {name}')

def clean(v):return str(v).strip() if v is not None else ''
def norm(v):return re.sub(r'\s+',' ',clean(v).replace('+',' ')).casefold()

p=Path('ciie-data.json')
prior=json.loads(p.read_text(encoding='utf8')) if p.exists() else {'buyers':[],'companies':{}}
old_buyers={clean(b.get('code')):b for b in prior['buyers']}
old_companies=prior['companies']
by_booth={(c.get('hall'),clean(c.get('booth'))):k for k,c in old_companies.items()}

buyers=[]
for row in list(wb['바이어'].values)[1:]:
 if not row or not clean(row[0]):continue
 code=clean(row[0]).upper(); b=old_buyers.get(code,{})
 buyers.append({'code':code,'name':clean(row[1]) or b.get('name',''),
   'cn':clean(row[2]),'meetings':b.get('meetings',[])})

companies={}
for row in list(wb['국내기업'].values)[1:]:
 if not row or not clean(row[1]):continue
 hall='consumer' if '소비' in clean(row[0]) else 'food'
 booth=clean(row[1]);oldkey=by_booth.get((hall,booth))
 key=oldkey or clean(row[2]) or f'{hall}-{booth}'
 c=dict(old_companies.get(oldkey,{})) if oldkey else {}
 c.update({'hall':hall,'booth':booth,'ko':clean(row[2]),'cn':clean(row[3]),'category':clean(row[4])})
 companies[key]=c

responses=[];seen=set()
for row in list(wb['Sheet1'].values)[1:]:
 if len(row)<7:continue
 buyer,seller=clean(row[5]),clean(row[6])
 if not buyer or not seller:continue
 k=(norm(buyer),norm(seller))
 if k in seen:continue
 seen.add(k);responses.append({'buyer':buyer,'seller':seller})

out={'updatedAt':datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC'),
     'buyers':buyers,'companies':companies,'responses':responses}
p.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(f'OK: {len(buyers)} buyers, {len(companies)} companies, {len(responses)} unique responses')
